import re
from pathlib import Path
from typing import Any
import docx
from backend.core.logging_config import ingestion_logger
from backend.core.models import ParsedBlock


def _parse_docx_file(file_path: Path) -> tuple[list[ParsedBlock], dict[str, Any]]:
    document = docx.Document(str(file_path))
    blocks: list[ParsedBlock] = []
    current_section = "Introduction"
    current_lines: list[str] = []
    section_index = 1

    def flush_section():
        nonlocal current_lines, section_index
        if current_lines:
            blocks.append(
                ParsedBlock(
                    text="\n".join(current_lines).strip(),
                    modality="docx",
                    page_number=section_index,
                    section_title=current_section,
                )
            )
            current_lines = []
            section_index += 1

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = (para.style.name or "").lower() if para.style else ""
        if "heading" in style_name or "title" in style_name:
            flush_section()
            current_section = text
            current_lines.append(f"## {text}")
        elif "list" in style_name:
            current_lines.append(f"- {text}")
        else:
            current_lines.append(text)

    flush_section()

    # Extract tables
    table_count = 0
    for idx, table in enumerate(document.tables, start=1):
        rows_data = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                rows_data.append(cells)
        if rows_data:
            table_count += 1
            header = rows_data[0]
            md_lines = [
                f"### Table {idx}",
                "| " + " | ".join(header) + " |",
                "| " + " | ".join(["---"] * len(header)) + " |",
            ]
            for r in rows_data[1:]:
                padded = r + [""] * max(0, len(header) - len(r))
                md_lines.append("| " + " | ".join(padded[: len(header)]) + " |")
            blocks.append(
                ParsedBlock(
                    text="\n".join(md_lines),
                    modality="docx",
                    page_number=section_index,
                    section_title=f"Table {idx}",
                    table_name=f"Table_{idx}",
                )
            )
            section_index += 1

    core_props = document.core_properties
    meta = {
        "author": getattr(core_props, "author", "") or "",
        "title": getattr(core_props, "title", "") or file_path.stem,
        "tables_extracted": table_count,
        "page_count": max(1, len(blocks)),
    }
    return blocks, meta


def _parse_legacy_doc_file(file_path: Path) -> tuple[list[ParsedBlock], dict[str, Any]]:
    """
    Parse legacy .doc or plain text without requiring Microsoft Office.
    Extracts UTF-8/Latin-1/UTF-16 printable text runs cleanly.
    """
    raw = file_path.read_bytes()
    # Try UTF-8 first (in case .doc is actually text/RTF/XML)
    try:
        decoded = raw.decode("utf-8")
        if sum(1 for c in decoded if c.isprintable() or c in "\r\n\t") / max(1, len(decoded)) > 0.85:
            text = decoded
        else:
            raise UnicodeDecodeError("utf-8", b"", 0, 1, "binary")
    except Exception:
        # Extract printable ASCII & UTF-16LE sequences from OLE2 binary .doc
        ascii_runs = re.findall(rb"[\x20-\x7E\r\n\t]{12,}", raw)
        utf16_runs = re.findall(rb"(?:[\x20-\x7E]\x00){12,}", raw)
        parts = [r.decode("latin-1", errors="ignore") for r in ascii_runs]
        parts += [r.decode("utf-16le", errors="ignore") for r in utf16_runs]
        text = "\n".join(parts)

    clean_text = re.sub(r"\n{3,}", "\n\n", text).strip()
    blocks = [
        ParsedBlock(
            text=clean_text,
            modality="doc",
            page_number=1,
            section_title=file_path.stem,
        )
    ]
    return blocks, {"page_count": 1, "title": file_path.stem}


def _parse_excel_file(file_path: Path) -> tuple[list[ParsedBlock], dict[str, Any]]:
    """
    Parse Excel (.xlsx) spreadsheets, serializing sheets into structured markdown tables.
    Preserves exact column headers, numerical values, and sheet names.
    """
    import openpyxl

    wb = openpyxl.load_workbook(str(file_path), data_only=True, read_only=True)
    blocks: list[ParsedBlock] = []
    sheet_names = wb.sheetnames
    table_count = 0

    for sheet_idx, sheet_name in enumerate(sheet_names, start=1):
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # Filter out empty rows
        non_empty_rows = [
            [str(c).strip() if c is not None else "" for c in r]
            for r in rows
            if any(c is not None and str(c).strip() != "" for c in r)
        ]
        if not non_empty_rows:
            continue

        headers = non_empty_rows[0]
        # Pad or clean headers
        clean_headers = [h if h else f"Col_{i+1}" for i, h in enumerate(headers)]
        md_lines = [
            f"### Sheet: {sheet_name}",
            "| " + " | ".join(clean_headers) + " |",
            "| " + " | ".join(["---"] * len(clean_headers)) + " |",
        ]
        for row in non_empty_rows[1:]:
            padded = row + [""] * max(0, len(clean_headers) - len(row))
            md_lines.append("| " + " | ".join(padded[: len(clean_headers)]) + " |")

        table_count += 1
        blocks.append(
            ParsedBlock(
                text="\n".join(md_lines),
                modality="text",
                page_number=sheet_idx,
                section_title=f"Sheet: {sheet_name}",
                table_name=sheet_name,
            )
        )

    wb.close()
    meta = {
        "sheets": sheet_names,
        "tables_extracted": table_count,
        "page_count": max(1, len(blocks)),
        "title": file_path.stem,
    }
    return blocks, meta


def parse_text_or_doc(file_path: Path) -> tuple[list[ParsedBlock], dict[str, Any]]:
    ext = file_path.suffix.lower()
    if ext == ".xlsx":
        blocks, meta = _parse_excel_file(file_path)
        ingestion_logger.info(f"Parsed Excel spreadsheet '{file_path.name}' with {len(blocks)} sheet tables.")
        return blocks, meta
    if ext == ".docx":
        blocks, meta = _parse_docx_file(file_path)
        ingestion_logger.info(f"Parsed DOCX '{file_path.name}' into {len(blocks)} structural blocks.")
        return blocks, meta
    if ext == ".doc":
        blocks, meta = _parse_legacy_doc_file(file_path)
        ingestion_logger.info(f"Parsed legacy DOC '{file_path.name}'.")
        return blocks, meta

    # Plain text (.txt, .md, .csv, .json, .xml)
    text = file_path.read_text(encoding="utf-8", errors="replace").strip()
    blocks = [
        ParsedBlock(
            text=text,
            modality="text",
            page_number=1,
            section_title=file_path.stem,
        )
    ]
    return blocks, {"page_count": 1, "title": file_path.stem}
